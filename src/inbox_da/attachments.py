"""Attachment parsing and Gemini vision routing."""

from __future__ import annotations

import base64
import csv
import json
from pathlib import Path
from typing import Any, Protocol

import fitz
from openpyxl import load_workbook

from inbox_da.config import Settings
from inbox_da.costing import collect_token_usage


class VisionParser(Protocol):
    def parse_image(self, path: Path, *, mime_type: str) -> dict[str, Any]: ...


class GeminiVisionParser:
    def __init__(self, model: Any) -> None:
        self.model = model

    def parse_image(self, path: Path, *, mime_type: str) -> dict[str, Any]:
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        content = [
            {
                "type": "text",
                "text": (
                    "Extract financial document text, key amounts, dates, counterparties, "
                    "and reference numbers from this attachment. Return concise JSON."
                ),
            },
            {"type": "image_url", "image_url": f"data:{mime_type};base64,{encoded}"},
        ]
        response = self.model.invoke(content)
        return {
            "model_response": getattr(response, "content", str(response)),
            "mime_type": mime_type,
            "token_usage": collect_token_usage(response).as_dict(),
        }


class OfflineVisionParser:
    """Deterministic local vision stand-in for offline demos and tests."""

    def parse_image(self, path: Path, *, mime_type: str) -> dict[str, Any]:
        return {
            "model_response": (
                f"Offline vision summary for {path.name}. The attachment appears to contain "
                "financial support data; use the email body and parsed attachments for amounts."
            ),
            "mime_type": mime_type,
            "offline": True,
            "token_usage": {
                "input_tokens": 0,
                "output_tokens": 0,
                "total_tokens": 0,
                "model_calls": 0,
            },
        }


def parse_attachment(
    path: Path,
    *,
    mime_type: str,
    settings: Settings,
    vision_parser: VisionParser | None = None,
) -> tuple[str, dict[str, Any], dict[str, Any] | None]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = list(csv.DictReader(path.read_text().splitlines()))
        text = json.dumps(rows, indent=2)
        return text, {"kind": "csv", "rows": rows}, None
    if suffix == ".txt" or mime_type.startswith("text/"):
        text = path.read_text()
        return text, {"kind": "text", "chars": len(text)}, None
    if suffix in {".xlsx", ".xlsm"}:
        rows = _read_xlsx(path)
        text = json.dumps(rows, indent=2)
        return text, {"kind": "xlsx", "rows": rows}, None
    if suffix == ".pdf" or mime_type == "application/pdf":
        text = _read_pdf_text(path)
        if len(text.strip()) >= settings.attachment_text_min_chars:
            return text, {"kind": "pdf_text", "chars": len(text)}, None
        summary = _require_vision(vision_parser).parse_image(path, mime_type=mime_type)
        return text, {"kind": "pdf_vision_fallback", "chars": len(text)}, summary
    if mime_type.startswith("image/") or suffix in {".png", ".jpg", ".jpeg"}:
        summary = _require_vision(vision_parser).parse_image(path, mime_type=mime_type)
        text = str(summary.get("model_response", ""))
        return text, {"kind": "image_vision"}, summary
    raise ValueError(f"Unsupported attachment type: {path} ({mime_type})")


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value) for value in rows[0]]
    return [dict(zip(headers, row, strict=False)) for row in rows[1:]]


def _read_pdf_text(path: Path) -> str:
    with fitz.open(path) as document:
        return "\n".join(page.get_text() for page in document)


def _require_vision(vision_parser: VisionParser | None) -> VisionParser:
    if vision_parser is None:
        raise RuntimeError("Vision parser is required for image or scanned PDF attachments")
    return vision_parser
